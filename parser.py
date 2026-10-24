"""Parse the faculty timetable workbook into schedule.json.
Nothing course-specific is hardcoded: every cell's text is read as-is.
Usage: python parser.py [schedule.xlsx] [schedule.json]
"""
import json, re, sys
import openpyxl
from openpyxl.utils import get_column_letter as L

FILE_NAME = sys.argv[1] if len(sys.argv) > 1 else "schedule.xlsx"
OUTPUT_FILE = sys.argv[2] if len(sys.argv) > 2 else "schedule.json"

DAYS = ["السبت", "الاحد", "الاثنين", "الثلاثاء", "الاربعاء", "الخميس"]
YEARS = {"الاول": 1, "الاولى": 1, "الثاني": 2, "الثانية": 2,
         "الثالث": 3, "الثالثة": 3, "الرابع": 4, "الرابعة": 4}
YEAR_NAMES = {1: "الفرقة الأولى", 2: "الفرقة الثانية", 3: "الفرقة الثالثة", 4: "الفرقة الرابعة"}
FIRST_HOUR, HOURS_PER_DAY, FIRST_COL = 8, 12, 3


def clean(v):
    return re.sub(r"\s+", " ", str(v)).strip() if v is not None else ""


def year_of(text):
    m = re.search(r"المستوى\s+(\S+)", clean(text))
    return YEARS.get(m.group(1).strip("-")) if m else None


def parse_text(text):
    """Return (course, instructor, location) from a cell's raw text."""
    tokens = []
    for line in str(text).split("\n"):
        tokens += [t.strip() for t in re.split(r"\s+-\s+", line) if t.strip()]
    course, instr, loc = [], "", ""
    for t in tokens:
        m = re.search(r"(?:^|\s)([أاإدم][\w.\s]{0,6}?\s*/)", t)   # م/ ... ، أ.د/ ... ، د/ ...
        if m and not instr:
            instr = clean(t[m.start(1):])
            t = t[:m.start(1)].strip()
            if not t:
                continue
        m = re.search(r"(مدرج|معمل|قاعة)\s*\S+", t)
        if m and not loc:
            loc = clean(m.group(0))
            t = (t[:m.start()] + t[m.end():]).strip()
        if t:
            course.append(clean(t))
    return " ".join(course), instr, loc


def numeric_cols(ws, row, c1, c2):
    out = {}
    for c in range(c1, c2 + 1):
        v = ws.cell(row, c).value
        if isinstance(v, (int, float)):
            out[c] = int(v)
    return out


def starts(ws, row, c1, c2):
    return [c for c in range(c1, c2 + 1)
            if clean(ws.cell(row, c).value) and not isinstance(ws.cell(row, c).value, (int, float))]


def merged_end(ws, row, col):
    for m in ws.merged_cells.ranges:
        if m.min_row <= row <= m.max_row and m.min_col == col:
            return m.max_col
    return col


def find_groups(ws):
    """Detect [(year, name, first_col, last_col, sec_row)] for a sheet."""
    maxc = ws.max_column
    # Layout C: tracks per level (levels in a row labelled 'المستوى' in col A)
    lvl_row = next((r for r in range(1, 15) if clean(ws.cell(r, 1).value) == "المستوى"), None)
    if lvl_row:
        trk_row, sec_row = lvl_row + 1, lvl_row + 2
        out = []
        for lc in starts(ws, lvl_row, FIRST_COL, maxc):
            word = clean(ws.cell(lvl_row, lc).value)
            if word not in YEARS:
                continue
            lend = merged_end(ws, lvl_row, lc)
            ts = [c for c in starts(ws, trk_row, lc, lend)]
            for i, tc in enumerate(ts):
                te = ts[i + 1] - 1 if i + 1 < len(ts) else lend
                out.append((YEARS[word], clean(ws.cell(trk_row, tc).value), tc, te, sec_row))
        return out
    year = year_of(ws["C2"].value)
    # Layout A: 'رقم المجموعة: N' blocks in row 3
    gs = [c for c in starts(ws, 3, FIRST_COL, maxc) if "المجموعة" in clean(ws.cell(3, c).value)]
    if gs:
        out = []
        for c in gs:
            num = re.search(r"\d+", clean(ws.cell(3, c).value))
            out.append((year, f"المجموعة {num.group(0)}" if num else clean(ws.cell(3, c).value),
                        c, merged_end(ws, 3, c), 5))
        return out
    # Layout B: department headers in row 4
    last = max(numeric_cols(ws, 5, FIRST_COL, maxc))
    ds = starts(ws, 4, FIRST_COL, maxc)
    return [(year, re.sub(r"\(\d+\)", "", clean(ws.cell(4, c).value)).strip(), c,
             (ds[i + 1] - 1 if i + 1 < len(ds) else last), 5) for i, c in enumerate(ds)]


def day_rows(ws):
    return [(r, clean(ws.cell(r, 1).value)) for r in range(1, ws.max_row + 1)
            if clean(ws.cell(r, 1).value) in DAYS]


def cell_blocks(ws, r1, r2):
    """Yield (min_row, max_row, min_col, max_col, value) for every non-empty cell,
    merged or not (the original parser skipped un-merged single-hour cells)."""
    merged = {}
    for m in ws.merged_cells.ranges:
        merged[(m.min_row, m.min_col)] = m
    for row in ws.iter_rows(min_row=r1, max_row=r2):
        for cell in row:
            if type(cell).__name__ == "MergedCell" or cell.value is None or not clean(cell.value):
                continue
            m = merged.get((cell.row, cell.column))
            yield (cell.row, m.max_row if m else cell.row, cell.column,
                   m.max_col if m else cell.column, cell.value)


years, events = {}, []
wb = openpyxl.load_workbook(FILE_NAME)
for ws in wb:
    groups = find_groups(ws)
    drs = day_rows(ws)
    if not groups or not drs:
        print(f"skip sheet {ws.title!r}")
        continue
    reg = []
    for year, name, c1, c2, sec_row in groups:
        secs = numeric_cols(ws, sec_row, c1, c2)
        y = years.setdefault(year, {"id": year, "name": YEAR_NAMES[year], "groups": []})
        gid = f"y{year}g{len(y['groups']) + 1}"
        y["groups"].append({"id": gid, "name": name, "sections": sorted(set(secs.values()))})
        reg.append((gid, secs))
    for (r, r2, c0, c2, val) in cell_blocks(ws, drs[0][0], drs[-1][0] + HOURS_PER_DAY - 1):
        day_start, day = max((d for d in drs if d[0] <= r), key=lambda d: d[0])
        course, instr, loc = parse_text(val)
        if not course or course.startswith("الدور"):   # floor labels, not classes
            continue
        for gid, secs in reg:
            hit = sorted(n for c, n in secs.items() if c0 <= c <= c2)
            if hit:
                events.append({
                    "group": gid, "day": DAYS.index(day),
                    "start": FIRST_HOUR + r - day_start,
                    "end": min(FIRST_HOUR + r2 - day_start + 1, 20),
                    "sections": hit, "course": course,
                    "instructor": instr, "location": loc})

events.sort(key=lambda e: (e["group"], e["day"], e["start"]))
out = {"years": [years[k] for k in sorted(years)], "events": events}
with open(OUTPUT_FILE, "w", encoding="utf-8") as f:
    json.dump(out, f, ensure_ascii=False, indent=1)
print(f"Extracted {len(events)} events -> {OUTPUT_FILE}")
